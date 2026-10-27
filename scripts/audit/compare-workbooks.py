"""Compares two workbooks the way a person would: is anything gone or different?

Usage: python3 compare-workbooks.py BEFORE.xlsx AFTER.xlsx [--ignore 'Sheet!A1' ...] [--layout-may-change]

Reads both with openpyxl (not the editor) and with a direct look at the package
parts. Prints JSON: {"same": bool, "differences": [...], "counts": {...}}.
Bytes, ids, the order of records and stored formula results are not compared;
content is. --ignore names cells that were edited on purpose.
--layout-may-change is for a file Microsoft Excel has re-saved: Excel adjusts
column widths and row heights to its own fonts, so those are not compared.
"""
import hashlib, json, re, sys, zipfile
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

args = sys.argv[1:]
ignored, layout, allowed, areas = set(), True, set(), []
while "--ignore" in args:
    at = args.index("--ignore"); ignored.add(args[at + 1]); del args[at:at + 2]
# --ignore-range 'Sheet!A1:C9': a block of cells that was edited on purpose.
while "--ignore-range" in args:
    at = args.index("--ignore-range")
    sheet_name, block = args[at + 1].rsplit("!", 1)
    low, high = (block.split(":") + [block])[:2]
    edge = lambda ref: (int(re.sub(r"[A-Z$]", "", ref)), sum((ord(c) - 64) * 26 ** i for i, c in enumerate(reversed(re.sub(r"[0-9$]", "", ref)))))
    areas.append((sheet_name, edge(low), edge(high))); del args[at:at + 2]
# --allow KIND: a kind of difference the action is meant to make (for example "merged").
while "--allow" in args:
    at = args.index("--allow"); allowed.add(args[at + 1]); del args[at:at + 2]
if "--layout-may-change" in args:
    args.remove("--layout-may-change"); layout = False
before_path, after_path = args[:2]
differences, counts = [], {}


def note(kind, where, was, now):
    if kind in allowed:
        return
    counts[kind] = counts.get(kind, 0) + 1
    if counts[kind] <= 3:
        differences.append({"kind": kind, "where": where, "before": str(was)[:160], "after": str(now)[:160]})


def content(cell):
    value = cell.value
    if isinstance(value, ArrayFormula):
        return ("array", value.ref, value.text)
    if isinstance(value, float):
        return ("number", round(value, 10))
    if isinstance(value, bool):
        return ("boolean", value)
    if isinstance(value, int):
        return ("number", float(value))
    if isinstance(value, str) and cell.data_type == "f":
        # '=A1' and '= A1' are the same formula; so are different quotings of a sheet name.
        return ("formula", re.sub(r"\s+", "", value).replace("'", ""))
    if value is None:
        return None
    return (type(value).__name__, str(value))


def colour(value):
    if value is None:
        return None
    if getattr(value, "type", None) == "rgb":
        return str(value.rgb)[-6:].upper() if value.rgb else None
    if getattr(value, "type", None) == "theme":
        return f"theme{value.theme}:{round(value.tint or 0, 3)}"
    if getattr(value, "type", None) == "indexed":
        return f"indexed{value.indexed}"
    return None


def look(cell):
    """How the cell is formatted, as a person would describe it."""
    if not cell.has_style:
        return None
    font, fill, align, border = cell.font, cell.fill, cell.alignment, cell.border
    described = (
        cell.number_format,
        bool(font.b), bool(font.i), font.u, font.sz, font.name, colour(font.color) or "theme1:0",
        fill.fill_type, colour(fill.fgColor) if fill.fill_type else None,
        align.horizontal, align.vertical, bool(align.wrap_text), align.indent or 0, align.text_rotation or 0,
        tuple(getattr(getattr(border, side), "style", None) for side in ("left", "right", "top", "bottom")),
        bool(cell.protection.locked), bool(cell.protection.hidden),
    )
    plain = ("General", False, False, None, 11.0, "Calibri", "theme1:0", None, None, None, None, False, 0, 0, (None,) * 4, True, False)
    return None if described == plain else described


def ordered_rules(sheet):
    """The rules in the order they apply. The order matters; the numbers that record it do not."""
    rules = sorted(
        ((rule.priority or 0, str(group.sqref), rule.type, rule.operator, tuple(rule.formula or ()))
         for group in sheet.conditional_formatting for rule in group.rules),
        key=lambda rule: (rule[0], rule[1:3]),
    )
    return [rule[1:] for rule in rules]


def sheet_facts(sheet):
    facts = {
        "merged": sorted(str(r) for r in sheet.merged_cells.ranges),
        "frozen": sheet.freeze_panes,
        "page breaks": (sorted(b.id for b in sheet.row_breaks.brk), sorted(b.id for b in sheet.col_breaks.brk)),
        "filter": sheet.auto_filter.ref,
        "filter columns": sorted((c.colId, bool(c.filters or c.customFilters)) for c in sheet.auto_filter.filterColumn),
        "filter buttons": sorted((c.colId, bool(c.hiddenButton), c.showButton is not False) for c in sheet.auto_filter.filterColumn),
        "sorted by": (sheet.auto_filter.sortState.ref, tuple((c.ref, bool(c.descending)) for c in sheet.auto_filter.sortState.sortCondition)) if sheet.auto_filter.sortState else None,
        "highlight rules": ordered_rules(sheet),
        "dropdowns": sorted((str(v.sqref), v.type, v.operator, v.formula1, v.formula2, bool(v.allow_blank)) for v in sheet.data_validations.dataValidation),
        "tables": sorted((t.name, t.ref, tuple(c.name for c in t.tableColumns), bool(t.totalsRowCount)) for t in sheet.tables.values()),
        "print area": sheet.print_area, "print titles": (sheet.print_title_rows, sheet.print_title_cols),
        "orientation": sheet.page_setup.orientation,
        "header and footer": (sheet.oddHeader.left.text, sheet.oddHeader.center.text, sheet.oddHeader.right.text,
                              sheet.oddFooter.left.text, sheet.oddFooter.center.text, sheet.oddFooter.right.text),
        "protected": bool(sheet.protection.sheet),
        "tab colour": colour(sheet.sheet_properties.tabColor),
        "hidden": sheet.sheet_state,
        "right to left": bool(sheet.sheet_view.rightToLeft),
        "gridlines": bool(sheet.sheet_view.showGridLines),
        "hidden columns": sorted(k for k, d in sheet.column_dimensions.items() if d.hidden),
        "hidden rows": sorted(k for k, d in sheet.row_dimensions.items() if d.hidden),
        "grouped columns": sorted((k, d.outlineLevel) for k, d in sheet.column_dimensions.items() if d.outlineLevel),
        "grouped rows": sorted((k, d.outlineLevel) for k, d in sheet.row_dimensions.items() if d.outlineLevel),
    }
    if layout:
        facts["column widths"] = sorted((k, round(d.width, 2)) for k, d in sheet.column_dimensions.items() if d.customWidth and d.width)
        facts["row heights"] = sorted((k, round(d.height, 2)) for k, d in sheet.row_dimensions.items() if d.height)
    return facts


def package_facts(path):
    facts = {}
    with zipfile.ZipFile(path) as package:
        names = package.namelist()
        def texts(pattern):
            return [package.read(n).decode("utf-8", "replace") for n in sorted(names) if re.match(pattern, n)]
        charts = texts(r"xl/charts/chart[^/]*\.xml$")
        facts["charts"] = len(charts)
        facts["chart data"] = sorted(ref.replace("'", "") for chart in charts for ref in re.findall(r"<(?:c:)?f>([^<]*)</(?:c:)?f>", chart))
        facts["chart kinds"] = sorted(kind for chart in charts for kind in re.findall(r"<c:(\w+Chart)>", chart))
        facts["pictures"] = sorted(hashlib.sha256(package.read(n)).hexdigest()[:16] for n in names if n.startswith("xl/media/"))
        facts["pivot tables"] = sum(bool(re.match(r"xl/pivotTables/[^/]+\.xml$", n)) for n in names)
        facts["pivot sources"] = sorted(m for text in texts(r"xl/pivotCache/[^/]+\.xml$") for m in re.findall(r"<worksheetSource\b[^>]*>", text))
        facts["links to other workbooks"] = sum(bool(re.match(r"xl/externalLinks/[^/]+\.xml$", n)) for n in names)
        facts["embedded files"] = sorted(hashlib.sha256(package.read(n)).hexdigest()[:16] for n in names if n.startswith("xl/embeddings/"))
        facts["macros"] = sorted(hashlib.sha256(package.read(n)).hexdigest()[:16] for n in names if n.endswith("vbaProject.bin"))
        facts["threaded comments"] = sum(len(re.findall(r"<threadedComment\b", t)) for t in texts(r"xl/threadedComments/[^/]+\.xml$"))
        facts["custom data"] = sorted(n for n in names if n.startswith("customXml/") and not n.endswith(".rels"))
        facts["shapes and text boxes"] = sum(len(re.findall(r"<xdr:sp\b", t)) for t in texts(r"xl/drawings/drawing[^/]*\.xml$"))
        facts["drawing anchors"] = sum(len(re.findall(r"<xdr:(?:twoCellAnchor|oneCellAnchor|absoluteAnchor)\b", t)) for t in texts(r"xl/drawings/drawing[^/]*\.xml$"))
        workbook = package.read("xl/workbook.xml").decode("utf-8", "replace")
        facts["1904 dates"] = bool(re.search(r'date1904="(1|true)"', workbook))
    return facts


try:
    before, after = load_workbook(before_path), load_workbook(after_path)
except Exception as error:  # a workbook the independent reader cannot open cannot be compared
    print(json.dumps({"same": None, "unreadable": str(error)[:200], "differences": [], "counts": {}}))
    sys.exit(0)

if before.sheetnames != after.sheetnames:
    note("sheets", "workbook", before.sheetnames, after.sheetnames)
names_before = {name: before.defined_names[name].attr_text.replace("'", "") for name in before.defined_names}
names_after = {name: after.defined_names[name].attr_text.replace("'", "") for name in after.defined_names}
if names_before != names_after:
    for name in sorted(set(names_before) | set(names_after)):
        if names_before.get(name) != names_after.get(name):
            note("names", name, names_before.get(name), names_after.get(name))

cells = 0
for name in before.sheetnames:
    if name not in after.sheetnames:
        continue
    a, b = before[name], after[name]
    if not hasattr(a, "iter_rows") or not hasattr(b, "iter_rows"):
        continue  # a chart sheet
    facts_a, facts_b = sheet_facts(a), sheet_facts(b)
    for key in facts_a:
        if facts_a[key] != facts_b[key]:
            note(key, name, facts_a[key], facts_b[key])
    seen = set()
    for source, other, direction in ((a, b, "before"), (b, a, "after")):
        for row in source.iter_rows():
            for cell in row:
                ref = cell.coordinate
                if ref in seen or f"{name}!{ref}" in ignored:
                    continue
                if any(area[0] == name and area[1][0] <= cell.row <= area[2][0] and area[1][1] <= cell.column <= area[2][1] for area in areas):
                    continue
                if cell.value is None and not cell.has_style and cell.comment is None and cell.hyperlink is None:
                    continue
                seen.add(ref); cells += 1
                first, second = (cell, other[ref]) if direction == "before" else (other[ref], cell)
                if content(first) != content(second):
                    note("cell contents", f"{name}!{ref}", content(first), content(second))
                elif look(first) != look(second):
                    note("cell formats", f"{name}!{ref}", look(first), look(second))
                link = lambda c: (c.hyperlink.target or c.hyperlink.location) if c.hyperlink else None
                if link(first) != link(second):
                    note("links", f"{name}!{ref}", link(first), link(second))
                text = lambda c: c.comment.text if c.comment else None
                if text(first) != text(second):
                    note("notes", f"{name}!{ref}", text(first), text(second))

package_before, package_after = package_facts(before_path), package_facts(after_path)
for key in package_before:
    if package_before[key] != package_after[key]:
        note(key, "workbook", package_before[key], package_after[key])

print(json.dumps({"same": not counts, "cells": cells, "differences": differences, "counts": counts}))
