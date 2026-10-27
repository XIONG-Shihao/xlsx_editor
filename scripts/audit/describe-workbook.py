"""Prints what an independent reader (openpyxl) sees in a saved workbook, as JSON.

Usage: python3 describe-workbook.py FILE [SHEET]
"""
import json, sys, zipfile, re
from openpyxl import load_workbook

book = load_workbook(sys.argv[1])
sheet = book[sys.argv[2]] if len(sys.argv) > 2 else book.worksheets[0]
charts = []
with zipfile.ZipFile(sys.argv[1]) as package:
    for name in sorted(n for n in package.namelist() if re.match(r"xl/charts/[^/]+\.xml$", n)):
        charts += re.findall(r"<(?:c:)?f>([^<]*)</(?:c:)?f>", package.read(name).decode())
# Checks openpyxl does not make but Excel does. Excel refuses or repairs a sheet
# whose records are out of the order the file format fixes, and a note whose
# box names a different cell from the note itself.
ORDER = """sheetPr dimension sheetViews sheetFormatPr cols sheetData sheetCalcPr sheetProtection
protectedRanges scenarios autoFilter sortState dataConsolidate customSheetViews mergeCells phoneticPr
conditionalFormatting dataValidations hyperlinks printOptions pageMargins pageSetup headerFooter
rowBreaks colBreaks customProperties cellWatches ignoredErrors smartTags drawing legacyDrawing
legacyDrawingHF drawingHF picture oleObjects controls webPublishItems tableParts extLst""".split()
problems = []


def column_number(letters):
    number = 0
    for letter in letters:
        number = number * 26 + ord(letter) - 64
    return number


def target(owner, value):
    if value.startswith("/"):
        return value[1:]
    folder = owner.rsplit("/", 1)[0].split("/")
    for step in value.split("/"):
        if step == "..":
            folder.pop()
        elif step != ".":
            folder.append(step)
    return "/".join(folder)


with zipfile.ZipFile(sys.argv[1]) as package:
    names = set(package.namelist())
    for name in sorted(n for n in names if re.match(r"xl/worksheets/[^/]+\.xml$", n)):
        text = package.read(name).decode()
        # Top-level records only: drop the cells, then read each record's name.
        outer = re.sub(r"<(\w+:)?sheetData[\s>].*?</(\w+:)?sheetData>", "<sheetData/>", text, flags=re.S)
        depth, seen = 0, []
        for closing, tag, selfclosing in re.findall(r"<(/?)(?:\w+:)?([A-Za-z0-9]+)(?:\s[^<>]*?)?(/?)>", outer):
            if closing:
                depth -= 1
                continue
            if depth == 1:
                seen.append(tag)
            if not selfclosing:
                depth += 1
        ranks = [ORDER.index(tag) for tag in seen if tag in ORDER]
        if ranks != sorted(ranks):
            problems.append(f"{name}: records out of order: {' '.join(seen)}")
        for tag in set(seen):
            if tag != "conditionalFormatting" and seen.count(tag) > 1:
                problems.append(f"{name}: {tag} appears {seen.count(tag)} times")
        rels = name.replace("worksheets/", "worksheets/_rels/") + ".rels"
        if rels not in names:
            continue
        links = re.findall(r"<Relationship\b[^>]*>", package.read(rels).decode())
        related = {}
        for link in links:
            kind = re.search(r'Type="[^"]*/([^"/]+)"', link)
            where = re.search(r'Target="([^"]*)"', link)
            if kind and where and "TargetMode=\"External\"" not in link:
                related.setdefault(kind.group(1), []).append(target(name, where.group(1)))
        for part in [p for group in related.values() for p in group]:
            if part not in names:
                problems.append(f"{name}: points at a missing part {part}")
        noted = set()
        for part in related.get("comments", []):
            if part in names:
                noted |= set(re.findall(r"<(?:\w+:)?comment\b[^>]*\bref=\"([A-Z]+[0-9]+)\"", package.read(part).decode()))
        boxes = set()
        for part in related.get("vmlDrawing", []):
            if part not in names:
                continue
            drawing = package.read(part).decode()
            for shape in re.findall(r"<(?:\w+:)?ClientData\b[^>]*ObjectType=\"Note\".*?</(?:\w+:)?ClientData>", drawing, flags=re.S):
                row = re.search(r"<(?:\w+:)?Row>\s*(\d+)\s*<", shape)
                column = re.search(r"<(?:\w+:)?Column>\s*(\d+)\s*<", shape)
                if row and column:
                    boxes.add((int(row.group(1)) + 1, int(column.group(1)) + 1))
        cells = set()
        for ref in noted:
            letters, digits = re.match(r"([A-Z]+)([0-9]+)", ref).groups()
            cells.add((int(digits), column_number(letters)))
        if cells != boxes:
            problems.append(f"{name}: notes are on {sorted(cells)} but their boxes are on {sorted(boxes)}")

print(json.dumps({
    "problems": problems,
    "sheets": book.sheetnames,
    "protected": bool(sheet.protection.sheet),
    "filter": sheet.auto_filter.ref,
    "filterButtons": sorted([c.colId, bool(c.hiddenButton)] for c in sheet.auto_filter.filterColumn if c.hiddenButton or c.showButton is False),
    "sortedBy": [[c.ref, bool(c.descending)] for c in sheet.auto_filter.sortState.sortCondition] if sheet.auto_filter.sortState else None,
    "filterColumns": [column.colId for column in sheet.auto_filter.filterColumn],
    "merged": sorted(str(r) for r in sheet.merged_cells.ranges),
    "frozen": sheet.freeze_panes,
    "breaks": {"rows": sorted(b.id for b in sheet.row_breaks.brk), "columns": sorted(b.id for b in sheet.col_breaks.brk)},
    "styled": {cell.coordinate: bool(cell.font.b) for row in sheet.iter_rows() for cell in row if cell.has_style and cell.font.b},
    "printArea": sheet.print_area,
    "rules": sorted(str(r.sqref) for r in sheet.conditional_formatting),
    "validation": sorted(str(v.sqref) for v in sheet.data_validations.dataValidation),
    "links": sorted(f"{cell.coordinate}={cell.hyperlink.target or cell.hyperlink.location}" for row in sheet.iter_rows() for cell in row if cell.hyperlink),
    "notes": sorted(f"{cell.coordinate}={cell.comment.text}" for row in sheet.iter_rows() for cell in row if cell.comment),
    "tables": sorted(f"{t.name}={t.ref}" for t in sheet.tables.values()),
    "names": sorted(f"{name}={book.defined_names[name].attr_text}" for name in book.defined_names),
    "charts": charts,
    "formulas": {cell.coordinate: cell.value for row in sheet.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")},
    "values": {cell.coordinate: cell.value for row in sheet.iter_rows(max_col=2) for cell in row if cell.value is not None},
    "thirdColumn": [cell.value for row in sheet.iter_rows(min_col=3, max_col=3) for cell in row],
}))
