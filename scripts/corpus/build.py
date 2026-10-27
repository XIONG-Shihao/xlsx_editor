"""Makes the part of the standing set that public files do not cover.

Usage: python3 scripts/corpus/build.py OUTPUT_FOLDER [PUBLIC_FOLDER]

Two kinds of workbook:
  everyday-*  workbooks of the kinds people keep (budget, invoice, timesheet ...),
              written by hand here, each using several features together;
  merged-*    real tables taken from the public workbooks, with features added
              that the public set lacks (charts, notes, dropdowns, grouping ...).

Everything is fixed: no random numbers, no clock. scripts/corpus/excel-save.mjs
can then pass these through Microsoft Excel so that they are Excel's own files.
"""
import datetime as dt, os, sys
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, ScatterChart, AreaChart, Reference, Series
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, IconSetRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableStyleInfo

out = sys.argv[1]
public = sys.argv[2] if len(sys.argv) > 2 else None
os.makedirs(out, exist_ok=True)
FIXED = dt.datetime(2026, 1, 5, 9, 0, 0)
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
made = []


def save(name, book):
    book.properties.created = FIXED
    book.properties.modified = FIXED
    book.properties.creator = "Workbook corpus"
    book.save(os.path.join(out, name + ".xlsx"))
    made.append(name)


def header(sheet, row, names, widths=None):
    for at, name in enumerate(names, 1):
        cell = sheet.cell(row=row, column=at, value=name)
        cell.font, cell.fill, cell.border = BOLD, HEAD, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if widths:
            sheet.column_dimensions[col(at)].width = widths[at - 1]


def lcg(seed):
    """A fixed sequence of numbers in [0, 1): the same on every run."""
    state = seed
    while True:
        state = (state * 1103515245 + 12345) % 2147483648
        yield state / 2147483648


# ---------------------------------------------------------------- everyday

def budget():
    book = Workbook(); s = book.active; s.title = "Budget 2026"
    s["A1"] = "Household budget 2026"; s["A1"].font = Font(bold=True, size=14); s.merge_cells("A1:N1")
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    header(s, 3, ["Category"] + months + ["Year"], [22] + [10] * 12 + [12])
    rows = [("Salary", 5200), ("Freelance", 600), ("Rent", -1650), ("Groceries", -540), ("Transport", -180),
            ("Utilities", -210), ("Insurance", -145), ("Leisure", -260), ("Savings", -800)]
    for r, (name, base) in enumerate(rows, 4):
        s.cell(row=r, column=1, value=name)
        for m in range(12):
            s.cell(row=r, column=2 + m, value=round(base * (1 + 0.01 * ((m * 7 + r) % 9 - 4)), 2)).number_format = '#,##0.00;[Red]-#,##0.00'
        s.cell(row=r, column=14, value=f"=SUM(B{r}:M{r})").number_format = '#,##0.00;[Red]-#,##0.00'
    last = 3 + len(rows)
    s.cell(row=last + 1, column=1, value="Balance").font = BOLD
    for c in range(2, 15):
        cell = s.cell(row=last + 1, column=c, value=f"=SUM({col(c)}4:{col(c)}{last})")
        cell.font = BOLD; cell.number_format = '#,##0.00;[Red]-#,##0.00'
    s.cell(row=last + 2, column=1, value="Savings rate")
    s.cell(row=last + 2, column=2, value=f"=-N{last}/SUM(N4:N5)").number_format = "0.0%"
    s.freeze_panes = "B4"
    s.conditional_formatting.add(f"B{last + 1}:N{last + 1}", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", bgColor="FFC7CE")))
    s.conditional_formatting.add(f"B4:M{last}", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    chart = LineChart(); chart.title = "Balance by month"; chart.height, chart.width = 7, 16
    chart.add_data(Reference(s, min_col=1, max_col=13, min_row=last + 1), from_rows=True, titles_from_data=True)
    chart.set_categories(Reference(s, min_col=2, max_col=13, min_row=3))
    s.add_chart(chart, f"A{last + 5}")
    s.print_area = f"A1:N{last + 2}"; s.page_setup.orientation = "landscape"; s.print_title_rows = "3:3"
    notes = book.create_sheet("Notes")
    notes["A1"] = "Assumptions"; notes["A1"].font = BOLD
    notes["A2"] = "Rent rises in July."; notes["A3"] = "Savings target"; notes["B3"] = 0.15; notes["B3"].number_format = "0%"
    book.defined_names["SavingsTarget"] = DefinedName("SavingsTarget", attr_text="Notes!$B$3")
    save("everyday-household-budget", book)


def invoice():
    book = Workbook(); s = book.active; s.title = "Invoice"
    s["A1"] = "INVOICE"; s["A1"].font = Font(bold=True, size=20); s.merge_cells("A1:C1")
    s["E1"] = "No."; s["F1"] = "INV-2026-0042"; s["E2"] = "Date"; s["F2"] = dt.date(2026, 3, 14); s["F2"].number_format = "yyyy-mm-dd"
    s["E3"] = "Due"; s["F3"] = "=F2+30"; s["F3"].number_format = "yyyy-mm-dd"
    s["A3"] = "Bill to"; s["A4"] = "Northwind Traders Ltd"; s["A5"] = "12 Harbour Road"; s["A6"] = "Dublin 2"
    header(s, 8, ["#", "Description", "Unit", "Quantity", "Unit price", "Amount"], [5, 38, 8, 10, 12, 14])
    items = [("Consulting, discovery phase", "day", 4, 780), ("Workshop facilitation", "day", 2, 950),
             ("Report writing", "hour", 14, 110), ("Travel", "trip", 3, 86.5), ("Printing", "set", 25, 4.2)]
    for r, (name, unit, qty, price) in enumerate(items, 9):
        s.cell(row=r, column=1, value=r - 8); s.cell(row=r, column=2, value=name); s.cell(row=r, column=3, value=unit)
        s.cell(row=r, column=4, value=qty); s.cell(row=r, column=5, value=price).number_format = "#,##0.00"
        s.cell(row=r, column=6, value=f"=ROUND(D{r}*E{r},2)").number_format = "#,##0.00"
        for c in range(1, 7): s.cell(row=r, column=c).border = BOX
    end = 8 + len(items)
    s[f"E{end + 1}"] = "Subtotal"; s[f"F{end + 1}"] = f"=SUM(F9:F{end})"
    s[f"E{end + 2}"] = "VAT"; s[f"D{end + 2}"] = 0.23; s[f"D{end + 2}"].number_format = "0%"; s[f"F{end + 2}"] = f"=ROUND(F{end + 1}*D{end + 2},2)"
    s[f"E{end + 3}"] = "Total"; s[f"F{end + 3}"] = f"=F{end + 1}+F{end + 2}"; s[f"F{end + 3}"].font = BOLD
    for r in range(end + 1, end + 4): s[f"F{r}"].number_format = '"€"#,##0.00'
    unit = DataValidation(type="list", formula1='"day,hour,trip,set,item"', allow_blank=True); unit.add(f"C9:C{end}"); s.add_data_validation(unit)
    qty = DataValidation(type="decimal", operator="greaterThan", formula1="0", showErrorMessage=True, errorTitle="Quantity", error="Enter a number above zero"); qty.add(f"D9:D{end}"); s.add_data_validation(qty)
    s[f"A{end + 5}"] = "Pay by bank transfer"; s[f"A{end + 6}"] = "Questions"; s[f"B{end + 6}"] = "accounts@example.com"; s[f"B{end + 6}"].hyperlink = "mailto:accounts@example.com"
    s["F1"].comment = Comment("Number is assigned by the accounts system.", "Accounts")
    s.print_area = f"A1:F{end + 6}"; s.oddFooter.center.text = "Page &P of &N"
    save("everyday-invoice", book)


def timesheet():
    book = Workbook(); s = book.active; s.title = "Week 11"
    header(s, 1, ["Date", "Day", "Start", "End", "Break (h)", "Hours", "Project", "Billable", "Notes"], [12, 6, 8, 8, 10, 8, 18, 9, 30])
    projects = ["Apollo", "Borealis", "Internal", "Apollo", "Cascade", "Apollo", "Borealis"]
    for i in range(7):
        r = 2 + i; day = dt.date(2026, 3, 9) + dt.timedelta(days=i)
        s.cell(row=r, column=1, value=day).number_format = "yyyy-mm-dd"
        s.cell(row=r, column=2, value=f'=TEXT(A{r},"ddd")')
        if i < 5:
            s.cell(row=r, column=3, value=dt.time(8, 30 + (i % 2) * 15)).number_format = "hh:mm"
            s.cell(row=r, column=4, value=dt.time(17, (i * 10) % 60)).number_format = "hh:mm"
            s.cell(row=r, column=5, value=0.5 + 0.25 * (i % 2))
        s.cell(row=r, column=6, value=f'=IF(OR(C{r}="",D{r}=""),0,ROUND((D{r}-C{r})*24-E{r},2))')
        s.cell(row=r, column=7, value=projects[i]); s.cell(row=r, column=8, value=projects[i] != "Internal")
    s["E9"] = "Total"; s["F9"] = "=SUM(F2:F8)"; s["E10"] = "Billable"; s["F10"] = '=SUMIF(H2:H8,TRUE,F2:F8)'
    s["E11"] = "Overtime"; s["F11"] = "=MAX(0,F9-40)"
    pick = DataValidation(type="list", formula1='"Apollo,Borealis,Cascade,Internal"'); pick.add("G2:G8"); s.add_data_validation(pick)
    s.conditional_formatting.add("A2:I8", FormulaRule(formula=['WEEKDAY($A2,2)>5'], fill=PatternFill("solid", bgColor="EDEDED")))
    s.auto_filter.ref = "A1:I8"; s.freeze_panes = "A2"
    s["I3"].comment = Comment("Client workshop ran late.", "Sam")
    summary = book.create_sheet("By project")
    header(summary, 1, ["Project", "Hours", "Share"], [16, 10, 10])
    for r, name in enumerate(["Apollo", "Borealis", "Cascade", "Internal"], 2):
        summary.cell(row=r, column=1, value=name)
        summary.cell(row=r, column=2, value=f"=SUMIF('Week 11'!G2:G8,A{r},'Week 11'!F2:F8)")
        summary.cell(row=r, column=3, value=f"=IF('Week 11'!$F$9=0,0,B{r}/'Week 11'!$F$9)").number_format = "0%"
    pie = PieChart(); pie.title = "Hours by project"
    pie.add_data(Reference(summary, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    pie.set_categories(Reference(summary, min_col=1, min_row=2, max_row=5)); summary.add_chart(pie, "E2")
    save("everyday-timesheet", book)


def inventory():
    book = Workbook(); s = book.active; s.title = "Stock"
    header(s, 1, ["SKU", "Item", "Category", "On hand", "Reorder at", "Unit cost", "Value", "Status", "Supplier"], [10, 26, 14, 9, 10, 10, 12, 12, 22])
    names = ["Bolt M8", "Nut M8", "Washer 8mm", "Bracket L", "Hinge 60", "Panel A", "Panel B", "Seal ring", "Cable 2m", "Fuse 5A", "Switch SP", "Lamp LED", "Tape 19mm", "Glue 50ml", "Filter F2", "Valve V1"]
    cats = ["Fasteners", "Fasteners", "Fasteners", "Hardware", "Hardware", "Panels", "Panels", "Seals", "Electrical", "Electrical", "Electrical", "Electrical", "Consumables", "Consumables", "Filters", "Valves"]
    rnd = lcg(11)
    for i, name in enumerate(names):
        r = 2 + i
        s.cell(row=r, column=1, value=f"SK-{1000 + i * 7}"); s.cell(row=r, column=2, value=name); s.cell(row=r, column=3, value=cats[i])
        s.cell(row=r, column=4, value=int(next(rnd) * 400)); s.cell(row=r, column=5, value=50 + (i % 4) * 25)
        s.cell(row=r, column=6, value=round(0.2 + next(rnd) * 30, 2)).number_format = "#,##0.00"
        s.cell(row=r, column=7, value=f"=D{r}*F{r}").number_format = "#,##0.00"
        s.cell(row=r, column=8, value=f'=IF(D{r}=0,"Out",IF(D{r}<=E{r},"Reorder","OK"))')
        s.cell(row=r, column=9, value=["Acme Supply", "Brightline", "Corvid Parts"][i % 3])
    end = 1 + len(names)
    table = Table(displayName="Stock", ref=f"A1:I{end}"); table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True); s.add_table(table)
    s.conditional_formatting.add(f"H2:H{end}", FormulaRule(formula=['$H2="Reorder"'], fill=PatternFill("solid", bgColor="FFEB9C")))
    s.conditional_formatting.add(f"H2:H{end}", FormulaRule(formula=['$H2="Out"'], fill=PatternFill("solid", bgColor="FFC7CE")))
    s.conditional_formatting.add(f"D2:D{end}", DataBarRule(start_type="min", end_type="max", color="5B9BD5"))
    s.freeze_panes = "C2"
    t = book.create_sheet("Totals")
    header(t, 1, ["Category", "Items", "Value"], [16, 8, 14])
    for r, name in enumerate(sorted(set(cats)), 2):
        t.cell(row=r, column=1, value=name)
        t.cell(row=r, column=2, value=f"=COUNTIF(Stock!C2:C{end},A{r})")
        t.cell(row=r, column=3, value=f"=SUMIF(Stock!C2:C{end},A{r},Stock!G2:G{end})").number_format = "#,##0.00"
    bar = BarChart(); bar.type = "bar"; bar.title = "Stock value by category"
    bar.add_data(Reference(t, min_col=3, min_row=1, max_row=1 + len(set(cats))), titles_from_data=True)
    bar.set_categories(Reference(t, min_col=1, min_row=2, max_row=1 + len(set(cats)))); t.add_chart(bar, "E2")
    save("everyday-inventory", book)


def gradebook():
    book = Workbook(); s = book.active; s.title = "Class 3B"
    header(s, 1, ["Student", "Quiz 1", "Quiz 2", "Midterm", "Project", "Final", "Weighted", "Grade", "Rank"], [20, 8, 8, 9, 9, 8, 10, 7, 6])
    weights = [0.1, 0.1, 0.25, 0.2, 0.35]
    s["A2"] = "Weight"; s["A2"].font = Font(italic=True)
    for c, w in enumerate(weights, 2): s.cell(row=2, column=c, value=w).number_format = "0%"
    names = ["Aoife Byrne", "Björn Ek", "Chloé Martin", "Daniel Ortiz", "Émile Roy", "Fatima Zahra", "Günter Weiß", "Hana Sato", "Ivan Petrov", "Jörg Müller", "Kaito Mori", "Léa Dubois", "Mário Silva", "Nora Åberg"]
    rnd = lcg(5)
    for i, name in enumerate(names):
        r = 3 + i; s.cell(row=r, column=1, value=name)
        for c in range(2, 7): s.cell(row=r, column=c, value=int(45 + next(rnd) * 55))
        s.cell(row=r, column=7, value=f"=SUMPRODUCT(B{r}:F{r},$B$2:$F$2)").number_format = "0.0"
        s.cell(row=r, column=8, value=f'=IF(G{r}>=85,"A",IF(G{r}>=70,"B",IF(G{r}>=55,"C","D")))')
        s.cell(row=r, column=9, value=f"=RANK(G{r},$G$3:$G${2 + len(names)})")
    end = 2 + len(names)
    for label, fn, r in [("Average", "AVERAGE", end + 1), ("Highest", "MAX", end + 2), ("Lowest", "MIN", end + 3), ("Median", "MEDIAN", end + 4)]:
        s.cell(row=r, column=1, value=label).font = BOLD
        for c in range(2, 8): s.cell(row=r, column=c, value=f"={fn}({col(c)}3:{col(c)}{end})").number_format = "0.0"
    s.conditional_formatting.add(f"B3:F{end}", IconSetRule("3TrafficLights1", "percent", [0, 55, 80]))
    s.conditional_formatting.add(f"H3:H{end}", CellIsRule(operator="equal", formula=['"D"'], font=Font(color="C00000", bold=True)))
    marks = DataValidation(type="whole", operator="between", formula1="0", formula2="100", showErrorMessage=True); marks.add(f"B3:F{end}"); s.add_data_validation(marks)
    s.freeze_panes = "B3"
    s.protection.sheet = True
    for r in range(3, end + 1):
        for c in range(2, 7): s.cell(row=r, column=c).protection = Protection(locked=False)
    save("everyday-gradebook", book)


def loan():
    book = Workbook(); s = book.active; s.title = "Loan"
    s["A1"] = "Loan amount"; s["B1"] = 240000; s["A2"] = "Annual rate"; s["B2"] = 0.0415; s["B2"].number_format = "0.00%"
    s["A3"] = "Years"; s["B3"] = 25; s["A4"] = "Payments per year"; s["B4"] = 12
    s["A5"] = "Payment"; s["B5"] = "=-PMT(B2/B4,B3*B4,B1)"; s["B5"].number_format = "#,##0.00"
    s["A6"] = "Total interest"; s["B6"] = "=B5*B3*B4-B1"; s["B6"].number_format = "#,##0.00"
    for name, ref in [("Principal", "$B$1"), ("Rate", "$B$2"), ("Years", "$B$3"), ("PerYear", "$B$4"), ("Payment", "$B$5")]:
        book.defined_names[name] = DefinedName(name, attr_text=f"Loan!{ref}")
    header(s, 8, ["Period", "Date", "Opening", "Interest", "Principal", "Closing"], [8, 12, 14, 12, 12, 14])
    for i in range(1, 61):
        r = 8 + i
        s.cell(row=r, column=1, value=i)
        s.cell(row=r, column=2, value=f"=EDATE(DATE(2026,1,1),A{r})").number_format = "mmm yyyy"
        s.cell(row=r, column=3, value="=Principal" if i == 1 else f"=F{r - 1}")
        s.cell(row=r, column=4, value=f"=C{r}*Rate/PerYear"); s.cell(row=r, column=5, value=f"=Payment-D{r}"); s.cell(row=r, column=6, value=f"=C{r}-E{r}")
        for c in range(3, 7): s.cell(row=r, column=c).number_format = "#,##0.00"
    for r in range(9, 69): s.row_dimensions[r].outlineLevel = 1
    s.freeze_panes = "A9"
    area = AreaChart(); area.title = "Balance, first five years"; area.add_data(Reference(s, min_col=6, min_row=8, max_row=68), titles_from_data=True); s.add_chart(area, "H2")
    c = book.create_sheet("Compare")
    header(c, 1, ["Rate", "10 years", "15 years", "20 years", "25 years", "30 years"], [10, 12, 12, 12, 12, 12])
    for i in range(8):
        r = 2 + i; c.cell(row=r, column=1, value=0.03 + i * 0.005).number_format = "0.0%"
        for k, years in enumerate([10, 15, 20, 25, 30], 2):
            c.cell(row=r, column=k, value=f"=-PMT($A{r}/12,{years}*12,Principal)").number_format = "#,##0"
    save("everyday-loan-schedule", book)


def schedule():
    book = Workbook(); s = book.active; s.title = "Plan"
    header(s, 1, ["Task", "Owner", "Start", "Days", "End", "Done", "Status"] + [f"W{w}" for w in range(1, 13)], [28, 10, 11, 6, 11, 7, 11] + [4] * 12)
    tasks = [("Requirements", "Ana", 0, 10, 1), ("Design", "Ben", 8, 15, 1), ("Prototype", "Cai", 20, 12, 0.6), ("Review", "Ana", 30, 5, 0),
             ("Build: core", "Dev", 34, 25, 0.2), ("Build: reports", "Eli", 40, 18, 0), ("Testing", "Fay", 58, 14, 0), ("Training", "Ben", 70, 6, 0), ("Launch", "Ana", 78, 2, 0)]
    for i, (name, owner, offset, days, done) in enumerate(tasks):
        r = 2 + i
        s.cell(row=r, column=1, value=name); s.cell(row=r, column=2, value=owner)
        s.cell(row=r, column=3, value=dt.date(2026, 2, 2) + dt.timedelta(days=offset)).number_format = "d mmm"
        s.cell(row=r, column=4, value=days); s.cell(row=r, column=5, value=f"=WORKDAY(C{r},D{r}-1)").number_format = "d mmm"
        s.cell(row=r, column=6, value=done).number_format = "0%"
        s.cell(row=r, column=7, value=f'=IF(F{r}>=1,"Done",IF(F{r}>0,"In progress","Not started"))')
        for w in range(12):
            s.cell(row=r, column=8 + w, value=f'=IF(AND($C{r}<=DATE(2026,2,2)+{w * 7 + 6},$E{r}>=DATE(2026,2,2)+{w * 7}),1,"")')
    end = 1 + len(tasks)
    s.conditional_formatting.add(f"H2:S{end}", CellIsRule(operator="equal", formula=["1"], fill=PatternFill("solid", bgColor="5B9BD5"), font=Font(color="5B9BD5")))
    s.conditional_formatting.add(f"F2:F{end}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="70AD47"))
    for c in range(8, 20): s.column_dimensions[col(c)].outlineLevel = 1
    s.freeze_panes = "B2"; s.auto_filter.ref = f"A1:G{end}"
    s["A6"].comment = Comment("Depends on the design review.", "Ana")
    s.sheet_properties.tabColor = "5B9BD5"
    save("everyday-project-schedule", book)


def sales_export():
    book = Workbook(); s = book.active; s.title = "orders_export"
    header(s, 1, ["order_id", "order_date", "customer", "country", "channel", "sku", "qty", "unit_price", "discount", "net", "paid"], [10, 12, 22, 12, 10, 10, 6, 10, 9, 12, 7])
    rnd = lcg(42)
    customers = ["Alder & Co", "Birch GmbH", "Cedar SARL", "Dogwood BV", "Elm Oy", "Fir AB", "Ginkgo KK", "Hazel Lda"]
    countries = ["Ireland", "Germany", "France", "Netherlands", "Finland", "Sweden", "Japan", "Portugal"]
    for i in range(600):
        r = 2 + i; k = int(next(rnd) * 8)
        s.cell(row=r, column=1, value=100200 + i); s.cell(row=r, column=2, value=dt.date(2025, 1, 1) + dt.timedelta(days=int(next(rnd) * 364))).number_format = "yyyy-mm-dd"
        s.cell(row=r, column=3, value=customers[k]); s.cell(row=r, column=4, value=countries[k]); s.cell(row=r, column=5, value=["web", "phone", "partner"][int(next(rnd) * 3)])
        s.cell(row=r, column=6, value=f"P-{int(next(rnd) * 40):03d}"); s.cell(row=r, column=7, value=1 + int(next(rnd) * 12))
        s.cell(row=r, column=8, value=round(5 + next(rnd) * 240, 2)); s.cell(row=r, column=9, value=[0, 0, 0.05, 0.1][int(next(rnd) * 4)]).number_format = "0%"
        s.cell(row=r, column=10, value=f"=ROUND(G{r}*H{r}*(1-I{r}),2)"); s.cell(row=r, column=11, value=next(rnd) > 0.2)
    s.auto_filter.ref = "A1:K601"; s.freeze_panes = "A2"
    p = book.create_sheet("by_country")
    header(p, 1, ["country", "orders", "net", "average", "share"], [14, 8, 14, 10, 8])
    for r, name in enumerate(countries, 2):
        p.cell(row=r, column=1, value=name); p.cell(row=r, column=2, value=f"=COUNTIF(orders_export!D:D,A{r})")
        p.cell(row=r, column=3, value=f"=SUMIF(orders_export!D:D,A{r},orders_export!J:J)").number_format = "#,##0.00"
        p.cell(row=r, column=4, value=f"=IF(B{r}=0,0,C{r}/B{r})").number_format = "#,##0.00"
        p.cell(row=r, column=5, value=f"=C{r}/SUM($C$2:$C$9)").number_format = "0.0%"
    p["A11"] = "largest"; p["B11"] = "=INDEX(A2:A9,MATCH(MAX(C2:C9),C2:C9,0))"
    p["A12"] = "lookup: Japan"; p["B12"] = '=VLOOKUP("Japan",A2:C9,3,FALSE)'
    bar = BarChart(); bar.title = "Net sales by country"; bar.add_data(Reference(p, min_col=3, min_row=1, max_row=9), titles_from_data=True)
    bar.set_categories(Reference(p, min_col=1, min_row=2, max_row=9)); p.add_chart(bar, "G2")
    save("everyday-sales-export", book)


def survey():
    book = Workbook(); s = book.active; s.title = "Responses"
    questions = ["Overall satisfaction", "Ease of use", "Value for money", "Support quality", "Would recommend"]
    header(s, 1, ["Respondent", "Age band", "Region"] + questions + ["Comment"], [12, 10, 10] + [12] * 5 + [36])
    rnd = lcg(77)
    comments = ["", "", "Fast delivery.", "Too expensive for what it is.", "", "Support was slow to answer — but helpful.", "", "Love it!", "Would like a dark theme; otherwise fine.", ""]
    for i in range(120):
        r = 2 + i; s.cell(row=r, column=1, value=f"R{i + 1:04d}")
        s.cell(row=r, column=2, value=["18-24", "25-34", "35-44", "45-54", "55+"][int(next(rnd) * 5)]); s.cell(row=r, column=3, value=["North", "South", "East", "West"][int(next(rnd) * 4)])
        for c in range(4, 9): s.cell(row=r, column=c, value=1 + int(next(rnd) * 5))
        s.cell(row=r, column=9, value=comments[i % len(comments)] or None)
    scale = DataValidation(type="whole", operator="between", formula1="1", formula2="5"); scale.add("D2:H121"); s.add_data_validation(scale)
    s.conditional_formatting.add("D2:H121", ColorScaleRule(start_type="num", start_value=1, start_color="F8696B", end_type="num", end_value=5, end_color="63BE7B"))
    s.auto_filter.ref = "A1:I121"; s.freeze_panes = "D2"
    a = book.create_sheet("Summary")
    header(a, 1, ["Question", "Mean", "Share 4-5", "1", "2", "3", "4", "5"], [24, 8, 10, 6, 6, 6, 6, 6])
    for r, q in enumerate(questions, 2):
        c = col(2 + r)
        a.cell(row=r, column=1, value=q); a.cell(row=r, column=2, value=f"=AVERAGE(Responses!{c}2:{c}121)").number_format = "0.00"
        a.cell(row=r, column=3, value=f'=COUNTIF(Responses!{c}2:{c}121,">=4")/COUNT(Responses!{c}2:{c}121)').number_format = "0%"
        for k in range(1, 6): a.cell(row=r, column=3 + k, value=f"=COUNTIF(Responses!{c}2:{c}121,{k})")
    stacked = BarChart(); stacked.type = "bar"; stacked.grouping = "percentStacked"; stacked.overlap = 100; stacked.title = "Answers"
    stacked.add_data(Reference(a, min_col=4, max_col=8, min_row=1, max_row=6), titles_from_data=True)
    stacked.set_categories(Reference(a, min_col=1, min_row=2, max_row=6)); a.add_chart(stacked, "A9")
    save("everyday-survey-results", book)


def bank_statement():
    book = Workbook(); s = book.active; s.title = "Statement"
    s["A1"] = "Account"; s["B1"] = "IE29 AIBK 9311 5212 3456 78"; s["A2"] = "Opening balance"; s["B2"] = 3120.55
    header(s, 4, ["Date", "Description", "Category", "Paid out", "Paid in", "Balance"], [12, 34, 14, 11, 11, 12])
    rnd = lcg(9)
    desc = [("SUPERVALU", "Groceries"), ("LEAP TOP-UP", "Transport"), ("ELECTRIC IRELAND", "Utilities"), ("SALARY ACME LTD", "Income"), ("CINEMA", "Leisure"), ("PHARMACY", "Health"), ("RENT STANDING ORDER", "Housing"), ("ATM WITHDRAWAL", "Cash")]
    for i in range(90):
        r = 5 + i; d, cat = desc[int(next(rnd) * len(desc))]
        s.cell(row=r, column=1, value=dt.date(2026, 1, 2) + dt.timedelta(days=i)).number_format = "dd/mm/yyyy"
        s.cell(row=r, column=2, value=d); s.cell(row=r, column=3, value=cat)
        amount = round(4 + next(rnd) * 180, 2)
        if cat == "Income": s.cell(row=r, column=5, value=round(2600 + next(rnd) * 200, 2))
        elif cat == "Housing": s.cell(row=r, column=4, value=1450)
        else: s.cell(row=r, column=4, value=amount)
        s.cell(row=r, column=6, value=f"=$B$2-SUM($D$5:D{r})+SUM($E$5:E{r})")
        for c in (4, 5, 6): s.cell(row=r, column=c).number_format = "#,##0.00"
    s.conditional_formatting.add("F5:F94", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000")))
    s.auto_filter.ref = "A4:F94"; s.freeze_panes = "A5"
    s.column_dimensions["C"].hidden = True
    s.row_dimensions[3].hidden = True
    save("everyday-bank-statement", book)


def contacts():
    book = Workbook(); s = book.active; s.title = "联系人"
    header(s, 1, ["姓名", "公司", "职位", "城市", "电话", "邮箱", "最近联系", "备注"], [12, 22, 14, 10, 16, 26, 12, 30])
    people = [("王芳", "睿创科技", "采购经理", "上海", "021-5555-0101"), ("李伟", "华信咨询", "总监", "北京", "010-5555-0144"), ("张敏", "南方电器", "工程师", "深圳", "0755-5555-0170"),
              ("刘洋", "东海贸易", "销售", "宁波", "0574-5555-0122"), ("陈静", "西部能源", "财务", "成都", "028-5555-0199"), ("赵磊", "北辰物流", "调度", "天津", "022-5555-0133")]
    for i, (name, firm, role, city, phone) in enumerate(people):
        r = 2 + i
        for c, v in enumerate([name, firm, role, city, phone], 1): s.cell(row=r, column=c, value=v)
        mail = f"contact{i + 1}@example.cn"; s.cell(row=r, column=6, value=mail).hyperlink = "mailto:" + mail
        s.cell(row=r, column=7, value=dt.date(2026, 2, 1) + dt.timedelta(days=i * 9)).number_format = 'yyyy"年"m"月"d"日"'
    s["H3"] = "下周回访"; s["H3"].comment = Comment("记得带样品", "销售部")
    s.auto_filter.ref = "A1:H7"; s.freeze_panes = "B2"
    t = book.create_sheet("统计")
    t["A1"] = "城市"; t["B1"] = "人数"
    for r, city in enumerate(["上海", "北京", "深圳", "宁波", "成都", "天津"], 2):
        t.cell(row=r, column=1, value=city); t.cell(row=r, column=2, value=f"=COUNTIF(联系人!D:D,A{r})")
    t["A9"] = "合计"; t["B9"] = "=SUM(B2:B7)"
    book.defined_names["联系人数"] = DefinedName("联系人数", attr_text="统计!$B$9")
    ar = book.create_sheet("ملخص")
    ar.sheet_view.rightToLeft = True
    ar["A1"] = "المدينة"; ar["B1"] = "العدد"; ar["A2"] = "الإجمالي"; ar["B2"] = "=统计!B9"
    jp = book.create_sheet("メモ 2026")
    jp["A1"] = "次回の打ち合わせ"; jp["B1"] = dt.date(2026, 4, 1); jp["B1"].number_format = "yyyy/mm/dd"; jp["A2"] = "参加者数"; jp["B2"] = "='联系人'!A1"
    save("everyday-contacts-multilingual", book)


def old_calendar():
    book = Workbook(); book.epoch = dt.datetime(1904, 1, 1)
    s = book.active; s.title = "Log"
    header(s, 1, ["Date", "Reading", "Change", "7-day average"], [12, 10, 10, 14])
    rnd = lcg(3)
    for i in range(60):
        r = 2 + i
        s.cell(row=r, column=1, value=dt.date(2026, 1, 1) + dt.timedelta(days=i)).number_format = "yyyy-mm-dd"
        s.cell(row=r, column=2, value=round(20 + next(rnd) * 6, 1))
        if i: s.cell(row=r, column=3, value=f"=B{r}-B{r - 1}")
        if i >= 6: s.cell(row=r, column=4, value=f"=AVERAGE(B{r - 6}:B{r})").number_format = "0.00"
    s["F1"] = "Days covered"; s["G1"] = "=MAX(A:A)-MIN(A:A)+1"
    s["F2"] = "First Monday"; s["G2"] = "=A2+MOD(8-WEEKDAY(A2,2),7)"; s["G2"].number_format = "yyyy-mm-dd"
    s["F3"] = "Month end"; s["G3"] = "=EOMONTH(A2,0)"; s["G3"].number_format = "yyyy-mm-dd"
    line = LineChart(); line.title = "Readings"; line.add_data(Reference(s, min_col=2, min_row=1, max_row=61), titles_from_data=True); s.add_chart(line, "F6")
    save("everyday-readings-1904-calendar", book)


def arrays_and_names():
    book = Workbook(); s = book.active; s.title = "Model"
    header(s, 1, ["Product", "Units", "Price", "Revenue"], [14, 8, 9, 12])
    data = [("Alpha", 120, 9.5), ("Beta", 80, 14.0), ("Gamma", 45, 31.25), ("Delta", 300, 2.2), ("Epsilon", 15, 120.0)]
    for i, (name, units, price) in enumerate(data):
        r = 2 + i; s.cell(row=r, column=1, value=name); s.cell(row=r, column=2, value=units); s.cell(row=r, column=3, value=price)
    s["D2"] = ArrayFormula("D2:D6", "=B2:B6*C2:C6")
    s["A8"] = "Total (array)"; s["D8"] = ArrayFormula("D8", "=SUM(B2:B6*C2:C6)")
    s["A9"] = "Total (SUMPRODUCT)"; s["D9"] = "=SUMPRODUCT(B2:B6,C2:C6)"
    s["A10"] = "Tax rate"; s["D10"] = 0.2; s["A11"] = "Tax"; s["D11"] = "=D9*TaxRate"
    s["A12"] = "Best seller"; s["D12"] = "=INDEX(Products,MATCH(MAX(Units),Units,0))"
    s["A13"] = "Text join"; s["D13"] = '=A2&" / "&A3&" ("&TEXT(C2,"0.00")&")"'
    s["A14"] = "Error on purpose"; s["D14"] = "=1/0"; s["A15"] = "Caught"; s["D15"] = '=IFERROR(D14,"n/a")'
    s["A16"] = "Nested"; s["D16"] = '=IF(AND(D9>1000,COUNTIF(Units,">100")>=2),"high","low")'
    for name, ref in [("Products", "Model!$A$2:$A$6"), ("Units", "Model!$B$2:$B$6"), ("TaxRate", "Model!$D$10")]:
        book.defined_names[name] = DefinedName(name, attr_text=ref)
    book.defined_names["Threshold"] = DefinedName("Threshold", attr_text="1000")
    scatter = ScatterChart(); scatter.title = "Units against price"; scatter.style = 13
    scatter.series.append(Series(Reference(s, min_col=2, min_row=2, max_row=6), Reference(s, min_col=3, min_row=2, max_row=6), title="Products")); s.add_chart(scatter, "F2")
    h = book.create_sheet("Hidden lookups"); h.sheet_state = "hidden"
    h["A1"] = "Code"; h["B1"] = "Label"
    for r, (code, label) in enumerate([("A", "Active"), ("P", "Paused"), ("X", "Closed")], 2): h.cell(row=r, column=1, value=code); h.cell(row=r, column=2, value=label)
    s["A18"] = "Status code"; s["B18"] = "P"; s["C18"] = "=VLOOKUP(B18,'Hidden lookups'!A2:B4,2,FALSE)"
    pick = DataValidation(type="list", formula1="'Hidden lookups'!$A$2:$A$4"); pick.add("B18"); s.add_data_validation(pick)
    save("everyday-model-arrays-names", book)


def wide_and_long():
    book = Workbook(); s = book.active; s.title = "Readings"
    header(s, 1, ["Station"] + [f"2026-{m:02d}" for m in range(1, 13)] * 5)
    rnd = lcg(123)
    for r in range(2, 402):
        s.cell(row=r, column=1, value=f"ST-{r - 1:04d}")
        for c in range(2, 62): s.cell(row=r, column=c, value=round(next(rnd) * 100, 2))
    s.cell(row=402, column=1, value="Mean").font = BOLD
    for c in range(2, 62): s.cell(row=402, column=c, value=f"=AVERAGE({col(c)}2:{col(c)}401)")
    s.freeze_panes = "B2"
    for n in range(2, 14):
        t = book.create_sheet(f"Month {n - 1:02d}")
        t["A1"] = "Station"; t["B1"] = "Value"; t["C1"] = "Against mean"
        for r in range(2, 42):
            t.cell(row=r, column=1, value=f"=Readings!A{r}"); t.cell(row=r, column=2, value=f"=Readings!{col(n)}{r}"); t.cell(row=r, column=3, value=f"=B{r}-Readings!{col(n)}$402")
    save("everyday-many-sheets-wide", book)


def expense_claim():
    book = Workbook(); s = book.active; s.title = "Claim"
    s.merge_cells("A1:G1"); s["A1"] = "Expense claim"; s["A1"].font = Font(bold=True, size=16); s["A1"].alignment = Alignment(horizontal="center")
    s["A2"] = "Employee"; s.merge_cells("B2:D2"); s["B2"] = "Nora Åberg"; s["E2"] = "Period"; s.merge_cells("F2:G2"); s["F2"] = "March 2026"
    s.merge_cells("A4:A5"); s["A4"] = "Date"; s.merge_cells("B4:B5"); s["B4"] = "Description"; s.merge_cells("C4:E4"); s["C4"] = "Amounts"
    s["C5"] = "Net"; s["D5"] = "VAT"; s["E5"] = "Gross"; s.merge_cells("F4:F5"); s["F4"] = "Category"; s.merge_cells("G4:G5"); s["G4"] = "Receipt"
    for cell in ["A4", "B4", "C4", "C5", "D5", "E5", "F4", "G4"]: s[cell].font, s[cell].fill, s[cell].alignment = BOLD, HEAD, Alignment(horizontal="center", vertical="center")
    rows = [(3, "Train to Cork", 58.0, "Travel"), (3, "Hotel, one night", 129.0, "Lodging"), (4, "Client dinner", 96.4, "Meals"), (9, "Taxi", 23.5, "Travel"), (17, "Conference fee", 350.0, "Training"), (24, "Stationery", 18.9, "Office")]
    for i, (day, text, net, cat) in enumerate(rows):
        r = 6 + i
        s.cell(row=r, column=1, value=dt.date(2026, 3, day)).number_format = "d mmm"; s.cell(row=r, column=2, value=text)
        s.cell(row=r, column=3, value=net); s.cell(row=r, column=4, value=f"=ROUND(C{r}*0.23,2)"); s.cell(row=r, column=5, value=f"=C{r}+D{r}")
        s.cell(row=r, column=6, value=cat); s.cell(row=r, column=7, value="yes" if i != 3 else "no")
        for c in (3, 4, 5): s.cell(row=r, column=c).number_format = "#,##0.00"
    end = 5 + len(rows)
    s[f"B{end + 1}"] = "Total"; s[f"B{end + 1}"].font = BOLD
    for c in ("C", "D", "E"): s[f"{c}{end + 1}"] = f"=SUM({c}6:{c}{end})"; s[f"{c}{end + 1}"].number_format = "#,##0.00"
    s[f"B{end + 2}"] = "Without receipt"; s[f"E{end + 2}"] = f'=SUMIF(G6:G{end},"no",E6:E{end})'
    cat = DataValidation(type="list", formula1='"Travel,Lodging,Meals,Training,Office"'); cat.add(f"F6:F{end}"); s.add_data_validation(cat)
    yes = DataValidation(type="list", formula1='"yes,no"'); yes.add(f"G6:G{end}"); s.add_data_validation(yes)
    s.conditional_formatting.add(f"G6:G{end}", CellIsRule(operator="equal", formula=['"no"'], fill=PatternFill("solid", bgColor="FFC7CE")))
    for c, w in zip("ABCDEFG", [10, 30, 10, 10, 10, 12, 9]): s.column_dimensions[c].width = w
    s[f"A{end + 4}"] = "Policy"; s[f"B{end + 4}"] = "Travel and expenses policy"; s[f"B{end + 4}"].hyperlink = "https://example.com/policy/expenses"
    save("everyday-expense-claim-merged-header", book)


def kpi_dashboard():
    book = Workbook(); d = book.active; d.title = "Data"
    header(d, 1, ["Month", "Revenue", "Cost", "Customers", "Churned", "NPS"], [10, 12, 12, 11, 9, 7])
    rnd = lcg(31)
    for i in range(24):
        r = 2 + i
        d.cell(row=r, column=1, value=dt.date(2024 + i // 12, i % 12 + 1, 1)).number_format = "mmm yy"
        d.cell(row=r, column=2, value=round(80000 + i * 2400 + next(rnd) * 9000)); d.cell(row=r, column=3, value=round(52000 + i * 1300 + next(rnd) * 6000))
        d.cell(row=r, column=4, value=1200 + i * 35 + int(next(rnd) * 40)); d.cell(row=r, column=5, value=int(10 + next(rnd) * 25)); d.cell(row=r, column=6, value=int(30 + next(rnd) * 35))
    s = book.create_sheet("Dashboard", 0)
    s["A1"] = "Company dashboard"; s["A1"].font = Font(bold=True, size=16); s.merge_cells("A1:H1")
    cards = [("Revenue, last month", "=INDEX(Data!B:B,COUNTA(Data!A:A))", "#,##0"), ("Margin", "=1-INDEX(Data!C:C,COUNTA(Data!A:A))/INDEX(Data!B:B,COUNTA(Data!A:A))", "0.0%"),
             ("Customers", "=INDEX(Data!D:D,COUNTA(Data!A:A))", "#,##0"), ("Churn", "=INDEX(Data!E:E,COUNTA(Data!A:A))/INDEX(Data!D:D,COUNTA(Data!A:A)-1)", "0.00%"),
             ("Revenue, 12 months", "=SUM(OFFSET(Data!B1,COUNTA(Data!A:A)-12,0,12,1))", "#,##0"), ("Growth, year on year", "=SUM(Data!B14:B25)/SUM(Data!B2:B13)-1", "0.0%")]
    for i, (label, formula, fmt) in enumerate(cards):
        r = 3 + (i // 3) * 3; c = 1 + (i % 3) * 3
        s.cell(row=r, column=c, value=label).font = Font(color="666666"); s.cell(row=r + 1, column=c, value=formula).number_format = fmt
        s.cell(row=r + 1, column=c).font = Font(bold=True, size=14)
    for c in "ADG": s.column_dimensions[c].width = 22
    line = LineChart(); line.title = "Revenue and cost"; line.height, line.width = 7, 15
    line.add_data(Reference(d, min_col=2, max_col=3, min_row=1, max_row=25), titles_from_data=True); line.set_categories(Reference(d, min_col=1, min_row=2, max_row=25)); s.add_chart(line, "A10")
    bars = BarChart(); bars.title = "NPS"; bars.height, bars.width = 7, 10
    bars.add_data(Reference(d, min_col=6, min_row=1, max_row=25), titles_from_data=True); s.add_chart(bars, "F10")
    s.sheet_view.showGridLines = False
    save("everyday-kpi-dashboard", book)


for build in [budget, invoice, timesheet, inventory, gradebook, loan, schedule, sales_export, survey, bank_statement, contacts,
              old_calendar, arrays_and_names, wide_and_long, expense_claim, kpi_dashboard]:
    build()

# ---------------------------------------------------------------- merged


def table_of(path, want_rows=12, max_rows=250, max_cols=10):
    """The first block of rows that looks like a table: a header row, then data."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in book.worksheets:
            rows = []
            for row in sheet.iter_rows(max_row=400, max_col=max_cols, values_only=True):
                rows.append(list(row))
            for start in range(len(rows) - want_rows):
                head = rows[start]
                width = sum(v is not None for v in head)
                if width < 3 or not all(isinstance(v, str) for v in head if v is not None):
                    continue
                cols_used = [i for i, v in enumerate(head) if v is not None]
                body = [[r[i] for i in cols_used] for r in rows[start + 1:start + 1 + max_rows]]
                body = [r for r in body if sum(v is not None for v in r) >= max(2, len(cols_used) // 2)]
                numeric = [i for i in range(len(cols_used)) if sum(isinstance(r[i], (int, float)) and not isinstance(r[i], bool) for r in body) > len(body) * 0.7]
                if len(body) >= want_rows and numeric:
                    return [str(head[i])[:40] for i in cols_used], body, numeric
    finally:
        book.close()
    return None


def clean(value):
    if isinstance(value, (int, float, str, bool, dt.date, dt.datetime, dt.time)) or value is None:
        return value[:500] if isinstance(value, str) else value
    return str(value)[:200]


def merged(name, source, kind):
    found = table_of(source)
    if not found:
        return False
    head, body, numeric = found
    book = Workbook(); s = book.active; s.title = "Data"
    s["A1"] = f"Source: {os.path.basename(source)[:60]}"; s["A1"].font = Font(italic=True, color="666666")
    top = 3
    header(s, top, head, [max(10, min(30, len(h) + 2)) for h in head])
    for r, row in enumerate(body, top + 1):
        for c, value in enumerate(row, 1): s.cell(row=r, column=c, value=clean(value))
    end = top + len(body); n = numeric[0] + 1; width = len(head)
    total = end + 1
    s.cell(row=total, column=1, value="Total").font = BOLD
    for i in numeric: s.cell(row=total, column=i + 1, value=f"=SUM({col(i + 1)}{top + 1}:{col(i + 1)}{end})").font = BOLD
    if kind == "chart-notes":
        chart = BarChart(); chart.title = head[numeric[0]][:30]
        chart.add_data(Reference(s, min_col=n, min_row=top, max_row=min(end, top + 20)), titles_from_data=True)
        chart.set_categories(Reference(s, min_col=1, min_row=top + 1, max_row=min(end, top + 20))); s.add_chart(chart, f"{col(width + 2)}3")
        s.cell(row=top + 1, column=n).comment = Comment("Checked against the published table.", "Reviewer")
        s.cell(row=top + 3, column=1).comment = Comment("Name as published.", "Reviewer")
        s.freeze_panes = s.cell(row=top + 1, column=2)
    elif kind == "table-rules":
        ok = [h if h else f"Column {i + 1}" for i, h in enumerate(head)]
        seen = {}
        for i, h in enumerate(ok):
            seen[h] = seen.get(h, 0) + 1
            if seen[h] > 1: ok[i] = f"{h} {seen[h]}"
            s.cell(row=top, column=i + 1, value=ok[i])
        s.cell(row=total, column=1, value=None)
        for i in numeric: s.cell(row=total, column=i + 1, value=None)
        table = Table(displayName="PublicData", ref=f"A{top}:{col(width)}{end}"); table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True); s.add_table(table)
        s.conditional_formatting.add(f"{col(n)}{top + 1}:{col(n)}{end}", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="5B9BD5"))
        s.cell(row=end + 3, column=1, value="Largest"); s.cell(row=end + 3, column=2, value=f"=MAX(PublicData[{ok[numeric[0]]}])")
        s.cell(row=end + 4, column=1, value="Rows"); s.cell(row=end + 4, column=2, value="=ROWS(PublicData)")
    elif kind == "filter-dropdown-group":
        s.auto_filter.ref = f"A{top}:{col(width)}{end}"
        flag = width + 1
        s.cell(row=top, column=flag, value="Reviewed").font = BOLD
        for r in range(top + 1, end + 1): s.cell(row=r, column=flag, value="no")
        pick = DataValidation(type="list", formula1='"yes,no,query"'); pick.add(f"{col(flag)}{top + 1}:{col(flag)}{end}"); s.add_data_validation(pick)
        for r in range(top + 1, end + 1): s.row_dimensions[r].outlineLevel = 1
        s.cell(row=total, column=flag, value=f'=COUNTIF({col(flag)}{top + 1}:{col(flag)}{end},"yes")')
        s.freeze_panes = s.cell(row=top + 1, column=1)
    elif kind == "second-sheet-names-links":
        book.defined_names["Figures"] = DefinedName("Figures", attr_text=f"Data!${col(n)}${top + 1}:${col(n)}${end}")
        book.defined_names["GrandTotal"] = DefinedName("GrandTotal", attr_text=f"Data!${col(n)}${total}")
        t = book.create_sheet("Résumé")
        t["A1"] = "Total"; t["B1"] = "=GrandTotal"; t["A2"] = "Mean"; t["B2"] = "=AVERAGE(Figures)"; t["A3"] = "Count"; t["B3"] = "=COUNT(Figures)"
        t["A4"] = "Top share"; t["B4"] = "=MAX(Figures)/GrandTotal"; t["B4"].number_format = "0.0%"
        t["A6"] = "Go to data"; t["A6"].hyperlink = f"#Data!A{top}"; t["A7"] = "Publisher"; t["A7"].hyperlink = "https://example.org/open-data"
        s.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(width, 6))
        line = LineChart(); line.title = "Series"; line.add_data(Reference(s, min_col=n, min_row=top, max_row=min(end, top + 30)), titles_from_data=True); t.add_chart(line, "D2")
    elif kind == "protected-print":
        s.protection.sheet = True
        for r in range(top + 1, end + 1): s.cell(row=r, column=n).protection = Protection(locked=False)
        s.print_area = f"A1:{col(width)}{total}"; s.print_title_rows = f"{top}:{top}"; s.page_setup.orientation = "landscape"; s.page_setup.fitToWidth = 1
        s.oddHeader.center.text = "Public data extract"; s.oddFooter.right.text = "&P / &N"
        s.column_dimensions[col(width)].hidden = width > 3
        s.conditional_formatting.add(f"{col(n)}{top + 1}:{col(n)}{end}", CellIsRule(operator="greaterThan", formula=[f"AVERAGE(${col(n)}${top + 1}:${col(n)}${end})"], font=Font(bold=True)))
    save(name, book)
    return True


if public:
    kinds = ["chart-notes", "table-rules", "filter-dropdown-group", "second-sheet-names-links", "protected-print"]
    count = 0
    for file in sorted(os.listdir(public)):
        if count >= 10 or not file.endswith(".xlsx") or file.startswith(("everyday-", "merged-", "~$")):
            continue
        if os.path.getsize(os.path.join(public, file)) > 600_000:
            continue
        try:
            kind = kinds[count % len(kinds)]
            if merged(f"merged-{kind}-{count + 1:02d}", os.path.join(public, file), kind):
                count += 1
        except Exception as error:
            print("skipped", file, str(error)[:80], file=sys.stderr)

print(f"made {len(made)} workbooks in {out}")
